from pathlib import Path

import openpyxl
import pytest

REAL_WORKBOOK = Path(__file__).resolve().parents[4] / "data" / "despensa_dona_maria.xlsx"
DESPENSA_HEADER = ("Ingrediente", "Quantidade em estoque", "Unidade")
PRECOS_HEADER = ("Ingrediente", "Quantidade comprada", "Unidade", "Preço total pago (R$)")


def real_rows(sheet: str) -> list[tuple]:
    workbook = openpyxl.load_workbook(REAL_WORKBOOK, read_only=True, data_only=True)
    _, *rows = workbook[sheet].iter_rows(values_only=True)
    return [row for row in rows if any(value is not None for value in row)]


@pytest.fixture
def make_workbook(tmp_path):
    """Build an .xlsx with the given rows; headers and sheet names can be broken on purpose."""

    def _make(despensa=(), precos=(), *, despensa_header=DESPENSA_HEADER, precos_header=PRECOS_HEADER,
              sheets=("Despensa", "Precos"), name="despensa.xlsx") -> Path:
        workbook = openpyxl.Workbook()
        workbook.remove(workbook.active)
        if "Despensa" in sheets:
            sheet = workbook.create_sheet("Despensa")
            sheet.append(despensa_header)
            for row in despensa:
                sheet.append(row)
        if "Precos" in sheets:
            sheet = workbook.create_sheet("Precos")
            sheet.append(precos_header)
            for row in precos:
                sheet.append(row)
        path = tmp_path / name
        workbook.save(path)
        return path

    return _make
