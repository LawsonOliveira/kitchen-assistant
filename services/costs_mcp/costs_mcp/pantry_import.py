"""Spreadsheet -> validated ingredient records (PLAN.md D31).

Strict on purpose: exact sheet/column names, names matched only after strip + NFC (fuzzy matching would
misprice silently), and every problem collected and raised together so the owner fixes the file once.
"""

import unicodedata
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from pathlib import Path

import openpyxl
import psycopg

from costs_mcp import db
from costs_mcp.pricing import format_brl
from costs_mcp.units import UnknownUnitError, parse_unit

SHEETS = {
    "Despensa": ("Ingrediente", "Quantidade em estoque", "Unidade"),
    "Precos": ("Ingrediente", "Quantidade comprada", "Unidade", "Preço total pago (R$)"),
}
CENT = Decimal("0.01")
DIFF_FIELDS = ("base_unit", "stock_base", "total_price_paid", "quantity_purchased_base")


class PantryImportError(ValueError):
    def __init__(self, errors: list[dict]):
        super().__init__("; ".join(error["message"] for error in errors))
        self.errors = errors


@dataclass(frozen=True)
class IngredientRecord:
    name: str
    base_unit: str
    stock_base: Decimal
    total_price_paid: Decimal
    quantity_purchased_base: Decimal
    purchase_unit_label: str


def _error(code: str, message: str, **details) -> dict:
    return {"code": code, "message": message, "details": details}


def _text(value) -> str:
    return unicodedata.normalize("NFC", str(value if value is not None else "")).strip()


def _decimal(value) -> Decimal | None:
    try:
        return Decimal(str(value)) if value is not None else None
    except InvalidOperation:
        return None


def plain(quantity: Decimal) -> str:
    """5 -> '5', 1.50 -> '1.5' (no exponent notation)."""
    return format(quantity.normalize(), "f")


def _read_sheets(path: Path, errors: list[dict]) -> dict[str, list[dict]]:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheets = {}
    for sheet, columns in SHEETS.items():
        if sheet not in workbook.sheetnames:
            errors.append(_error("missing_sheet", f"sheet {sheet!r} is missing", sheet=sheet))
            continue
        header, *rows = list(workbook[sheet].iter_rows(values_only=True)) or [()]
        header = [_text(cell) for cell in header]
        missing = [column for column in columns if column not in header]
        if missing:
            errors.append(_error("missing_column", f"sheet {sheet!r} lacks columns {missing}", sheet=sheet, columns=missing))
            continue
        positions = {column: header.index(column) for column in columns}
        sheets[sheet] = [{c: row[i] for c, i in positions.items()} for row in rows if any(v is not None for v in row)]
    return sheets


def _record(name: str, stock_row: dict, price_row: dict, errors: list[dict]) -> IngredientRecord | None:
    stock_unit, price_unit = _text(stock_row["Unidade"]), _text(price_row["Unidade"])
    try:
        stock_spec, price_spec = parse_unit(stock_unit), parse_unit(price_unit)
    except UnknownUnitError as error:
        errors.append(_error("unknown_unit", f"{name!r}: unknown unit {error.raw!r}", ingredient=name, unit=error.raw))
        return None
    if stock_spec.base_unit != price_spec.base_unit:
        errors.append(_error("incompatible_units", f"{name!r}: stock in {stock_unit} but bought in {price_unit}",
                             ingredient=name, stock_unit=stock_unit, price_unit=price_unit))
        return None
    row_errors = []
    stock, purchased = _decimal(stock_row["Quantidade em estoque"]), _decimal(price_row["Quantidade comprada"])
    if stock is None or stock < 0:
        row_errors.append(_error("non_positive_quantity", f"{name!r}: invalid stock {stock_row['Quantidade em estoque']}", ingredient=name))
    if purchased is None or purchased <= 0:
        row_errors.append(_error("non_positive_quantity", f"{name!r}: invalid quantity purchased {price_row['Quantidade comprada']}", ingredient=name))
    total = _decimal(price_row["Preço total pago (R$)"])
    if total is None or total <= 0:
        row_errors.append(_error("invalid_price", f"{name!r}: invalid price {price_row['Preço total pago (R$)']}", ingredient=name))
    elif abs(total - total.quantize(CENT)) > Decimal("0.000001"):  # float artifacts like 79.90000000000001 pass
        row_errors.append(_error("non_cent_price", f"{name!r}: price {total} is not a whole number of cents", ingredient=name))
    errors.extend(row_errors)
    if row_errors:
        return None
    return IngredientRecord(name, stock_spec.base_unit, stock * stock_spec.factor_to_base, total.quantize(CENT),
                            purchased * price_spec.factor_to_base, f"{plain(purchased)} {price_unit}")


def load_records(path: Path | str) -> list[IngredientRecord]:
    errors: list[dict] = []
    sheets = _read_sheets(Path(path), errors)
    if errors:
        raise PantryImportError(errors)
    pantry = {_text(row["Ingrediente"]): row for row in sheets["Despensa"]}
    prices = {_text(row["Ingrediente"]): row for row in sheets["Precos"]}
    for name in sorted(pantry.keys() ^ prices.keys()):
        errors.append(_error("name_in_one_sheet", f"{name!r} appears in only one sheet", ingredient=name))
    records = [record for name in pantry if name in prices
               if (record := _record(name, pantry[name], prices[name], errors)) is not None]
    if errors:
        raise PantryImportError(errors)
    return records


def diff(current: list[IngredientRecord], new: list[IngredientRecord]) -> dict:
    old_by_name = {record.name: record for record in current}
    new_by_name = {record.name: record for record in new}
    changed = []
    for name in sorted(old_by_name.keys() & new_by_name.keys()):
        old, fresh = old_by_name[name], new_by_name[name]
        fields = {field: [_show(field, getattr(old, field)), _show(field, getattr(fresh, field))]
                  for field in DIFF_FIELDS if getattr(old, field) != getattr(fresh, field)}
        if fields:
            changed.append({"name": name, "fields": fields, "display": _change_display(name, fields, fresh.base_unit)})
    return {"added": sorted(new_by_name.keys() - old_by_name.keys()),
            "removed": sorted(old_by_name.keys() - new_by_name.keys()), "changed": changed}


FIELD_LABELS = {"base_unit": "unidade", "stock_base": "estoque", "total_price_paid": "preço pago",
                "quantity_purchased_base": "quantidade comprada"}


def _change_display(name: str, fields: dict, base_unit: str) -> str:
    """One owner-facing line per changed ingredient. Money is formatted here because Dona Sálvia may only repeat amounts
    that came from a display string (D12)."""
    parts = []
    for field, (old_value, new_value) in fields.items():
        if field == "total_price_paid":
            old_value, new_value = format_brl(Decimal(old_value)), format_brl(Decimal(new_value))
        elif field in ("stock_base", "quantity_purchased_base"):
            old_value, new_value = f"{old_value} {base_unit}", f"{new_value} {base_unit}"
        parts.append(f"{FIELD_LABELS[field]} de {old_value} para {new_value}")
    return f"{name}: {'; '.join(parts)}"


def _show(field: str, value) -> str:
    if field == "total_price_paid":
        return str(value.quantize(CENT))
    return plain(value) if isinstance(value, Decimal) else str(value)


def seed_from_workbook(conn: psycopg.Connection, path: Path) -> int:
    """Seed an empty database; an invalid workbook raises, so the server fails to start (fail loud)."""
    records = load_records(path)
    if db.count_ingredients(conn):
        return 0
    for record in records:
        ingredient_id = db.insert_ingredient(conn, record.name, record.base_unit)
        db.insert_pantry_stock(conn, ingredient_id, record.stock_base)
        db.insert_price(conn, ingredient_id, record.total_price_paid, record.quantity_purchased_base,
                        record.purchase_unit_label, "spreadsheet")
    conn.commit()
    return len(records)
